#!/usr/bin/env python3
"""Bounded local input preparation; never calls a model, network or database."""
import argparse
import csv
import datetime as dt
import json
import math
import re
import sys
import zipfile
from pathlib import Path

from validate import load_json, validate_config, validate_input

MAX_BYTES = 8 * 1024 * 1024
MAX_ROWS = 10000


def pointer_escape(value):
    return str(value).replace('~', '~0').replace('/', '~1')


def resolve(data, pointer):
    if pointer == '':
        return data
    if not pointer.startswith('/') or re.search(r'~(?![01])', pointer):
        raise ValueError('Invalid JSON pointer')
    for token in pointer[1:].split('/'):
        token = token.replace('~1', '/').replace('~0', '~')
        if isinstance(data, list):
            if not re.fullmatch(r'0|[1-9][0-9]*', token):
                raise ValueError('Invalid array index')
            data = data[int(token)]
        else:
            data = data[token]
    return data


def bounds(value, allow_zero=False):
    match = re.fullmatch(r'(\d+):(\d+)', value or '')
    if not match:
        raise ValueError('Range must be START:END inclusive')
    start, end = map(int, match.groups())
    if start < (0 if allow_zero else 1) or end < start or end - start + 1 > MAX_ROWS:
        raise ValueError('Invalid or excessive range')
    return start, end


def convert(value, col):
    """Only explicit conversions. Empty string, None, zero and False differ."""
    if value is None:
        return None
    if value == '' and isinstance(value, str):
        return None if col['empty'] == 'null' else ''
    if isinstance(value, str) and value in col.get('null_tokens', []):
        return None
    kind = col['type']
    if kind == 'string':
        if not isinstance(value, str):
            raise ValueError(f"{col['source']}: expected a string; numeric/bool/date coercion is not implicit")
        return value
    if kind in ('number', 'integer'):
        if isinstance(value, bool):
            raise ValueError(f"{col['source']}: boolean is not a number")
        if isinstance(value, str):
            if not re.fullmatch(r'-?(?:0|[1-9][0-9]*)(?:\.[0-9]+)?(?:[eE][+-]?[0-9]+)?', value):
                raise ValueError(f"{col['source']}: ambiguous numeric format")
            value = json.loads(value)
        if not isinstance(value, (int, float)) or (isinstance(value, float) and not math.isfinite(value)):
            raise ValueError(f"{col['source']}: invalid number")
        if kind == 'integer' and (isinstance(value, float) and not value.is_integer()):
            raise ValueError(f"{col['source']}: expected integer")
        return int(value) if kind == 'integer' else value
    if kind == 'boolean':
        if isinstance(value, bool):
            return value
        if isinstance(value, str):
            if value in col.get('true_values', []):
                return True
            if value in col.get('false_values', []):
                return False
        raise ValueError(f"{col['source']}: boolean needs explicit true_values/false_values")
    if kind in ('date', 'datetime'):
        if isinstance(value, str):
            if not col.get('date_format'):
                raise ValueError(f"{col['source']}: date_format required")
            value = dt.datetime.strptime(value, col['date_format'])
        if kind == 'date' and isinstance(value, (dt.date, dt.datetime)):
            return value.date().isoformat() if isinstance(value, dt.datetime) else value.isoformat()
        if kind == 'datetime' and isinstance(value, dt.datetime):
            if value.tzinfo is None:
                raise ValueError(f"{col['source']}: datetime timezone is missing")
            return value.isoformat()
        raise ValueError(f"{col['source']}: wrong date type")
    raise ValueError('Unsupported mapping type')


def mapped_rows(rows, headers, cols, physical):
    if len(set(headers)) != len(headers):
        raise ValueError('Duplicate headers')
    if any(not isinstance(h, str) or not h for h in headers):
        raise ValueError('Headers must be nonempty strings')
    if not cols:
        raise ValueError('CSV/XLSX require explicit config.columns selection and type mapping')
    indices = []
    for col in cols:
        if col['source'] not in headers:
            raise ValueError(f"Missing column: {col['source']}")
        indices.append(headers.index(col['source']))
    result, locations = [], {}
    for i, row in enumerate(rows):
        if len(row) != len(headers):
            raise ValueError(f'Ragged record {i + 1}')
        item = {}
        for col, index in zip(cols, indices):
            item[col['field']] = convert(row[index], col)
            locations[f'/{i}/{pointer_escape(col["field"])}'] = physical(i, index)
        result.append(item)
    if not result:
        raise ValueError('Requested range is empty')
    return result, locations


def read_csv(args, cols):
    start, end = bounds(args.rows)
    if len(args.delimiter) != 1:
        raise ValueError('CSV delimiter must be one character')
    rows = []
    with args.file.open(encoding=args.encoding, newline='') as file:
        reader = csv.reader(file, delimiter=args.delimiter, strict=True)
        headers = next(reader, None)
        if headers is None:
            raise ValueError('CSV has no header')
        for number, row in enumerate(reader, 1):
            if number >= start:
                rows.append(row)
            if number == end:
                break
    if len(rows) != end - start + 1:
        raise ValueError('Requested CSV records do not all exist; narrow the range explicitly')
    state, locations = mapped_rows(rows, headers, cols, lambda i, j: f'data record {start+i}, column {headers[j]} (header record 1)')
    return state, locations, f'CSV selected data records {start}:{end}; columns {[c["source"] for c in cols]}; header and preceding records scanned; later records not inspected'


def read_xlsx(args, cols):
    if not args.sheet or not args.cell_range:
        raise ValueError('XLSX requires --sheet and --range including header row')
    if not re.fullmatch(r'[A-Z]+[1-9][0-9]*:[A-Z]+[1-9][0-9]*', args.cell_range):
        raise ValueError('XLSX range must be explicit A1:B3 form')
    import openpyxl
    from openpyxl.utils.cell import range_boundaries, get_column_letter
    # Local archive resource guard, no external links or macro execution.
    with zipfile.ZipFile(args.file) as archive:
        if sum(i.file_size for i in archive.infolist()) > 64 * 1024 * 1024:
            raise ValueError('XLSX uncompressed size exceeds 64 MiB; provide a bounded export')
    min_col, min_row, max_col, max_row = range_boundaries(args.cell_range)
    if max_row <= min_row or max_col < min_col or (max_row-min_row)*(max_col-min_col+1) > MAX_ROWS:
        raise ValueError('Invalid or excessive XLSX range (maximum 10000 data cells)')
    wb = openpyxl.load_workbook(args.file, read_only=True, data_only=False, keep_links=False)
    try:
        if args.sheet not in wb.sheetnames:
            raise ValueError('Requested sheet does not exist')
        sheet = wb[args.sheet]
        if sheet.max_row is not None and (max_row > sheet.max_row or max_col > sheet.max_column):
            raise ValueError('Requested range exceeds worksheet dimensions')
        raw = list(sheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col))
        headers = [cell.value for cell in raw[0]]
        selected = {c['source'] for c in cols}
        indices = {i for i, h in enumerate(headers) if h in selected}
        for row in raw[1:]:
            for i in indices:
                if row[i].data_type in ('f', 'e'):
                    raise ValueError(f'Formula/error at {row[i].coordinate}; provide verified values, not cached guesses')
        values = [['' if cell.data_type == 'inlineStr' and cell.value is None else cell.value for cell in row] for row in raw[1:]]
        state, locations = mapped_rows(values, headers, cols, lambda i, j: f'{args.sheet}!{get_column_letter(min_col+j)}{min_row+1+i}')
        return state, locations, f'XLSX {args.sheet}!{args.cell_range}; first row is header; selected columns {[c["source"] for c in cols]}; outside range not inspected'
    finally:
        wb.close()


def read_text(args):
    start, end = bounds(args.lines)
    selected = []
    with args.file.open(encoding=args.encoding) as file:
        for number, line in enumerate(file, 1):
            if number >= start:
                selected.append(line)
            if number == end:
                break
    if len(selected) != end-start+1:
        raise ValueError('Requested text lines do not all exist')
    return ''.join(selected), {'': f'lines {start}:{end}'}, f'text lines {start}:{end}; other lines not inspected'


def make_units(state, config, source_id, complete):
    opts = config['unit']
    mode = opts['mode']
    if mode in ('single', 'whole'):
        return state, [{'subject_id': opts.get('subject_id', 'subject-1'), 'pointer': '', 'source_ids': [source_id]}]
    if not isinstance(state, list) or not state:
        raise ValueError('row/group mode needs a nonempty JSON array')
    if mode == 'row':
        if not opts.get('id_field'):
            raise ValueError('row mode requires unit.id_field; do not invent business identifiers')
        subjects = []
        seen = set()
        for i, row in enumerate(state):
            if not isinstance(row, dict) or opts['id_field'] not in row:
                raise ValueError('id_field missing in row')
            value = row[opts['id_field']]
            if not isinstance(value, str) or not value or value in seen or value == '__request__':
                raise ValueError('Row IDs must be distinct nonempty strings, not reserved IDs')
            seen.add(value)
            subjects.append({'subject_id': value, 'pointer': f'/{i}', 'source_ids': [source_id]})
        return state, subjects
    keys = opts.get('group_by')
    if not keys:
        raise ValueError('group mode requires group_by')
    groups, lookup = [], {}
    for i, row in enumerate(state):
        if not isinstance(row, dict) or any(k not in row or row[k] is None or row[k] == '' for k in keys):
            raise ValueError('Missing group key')
        if any(isinstance(row[k], (list, dict)) for k in keys):
            raise ValueError('Group keys must be scalar')
        key = json.dumps([row[k] for k in keys], ensure_ascii=False, separators=(',', ':'), allow_nan=False)
        if key not in lookup:
            lookup[key] = len(groups)
            groups.append({'key': {k:row[k] for k in keys}, 'rows': [], 'input_indices': []})
        group = groups[lookup[key]]
        group['rows'].append(row)
        group['input_indices'].append(i)
    # IDs explicitly derived from exact key tuples, never guessed entity IDs.
    subjects = [{'subject_id': 'group:'+key, 'pointer': f'/{index}', 'source_ids':[source_id]} for key, index in lookup.items()]
    return groups, subjects


def prepare(args):
    config = load_json(args.config)
    errors = validate_config(config)
    if errors:
        raise ValueError('; '.join(errors))
    if args.file.stat().st_size > MAX_BYTES:
        raise ValueError('Input exceeds 8 MiB; supply a bounded export/chunk')
    if args.format == 'csv':
        state, locations, read_scope = read_csv(args, config.get('columns', []))
    elif args.format == 'xlsx':
        state, locations, read_scope = read_xlsx(args, config.get('columns', []))
    elif args.format in ('text', 'document'):
        state, locations, read_scope = read_text(args)
    else:
        if not args.whole_file:
            raise ValueError('JSON parser reads the entire file; --whole-file explicitly authorizes this bounded file read')
        state = resolve(load_json(args.file), args.pointer)
        locations = {'': f'JSON pointer {args.pointer or "(root)"}'}
        read_scope = f'Whole JSON file parsed (explicit --whole-file); selected pointer {args.pointer or "(root)"}'
        if args.items:
            if not isinstance(state, list):
                raise ValueError('--items requires an array')
            start, end = bounds(args.items, allow_zero=True)
            if end >= len(state):
                raise ValueError('Requested JSON items do not all exist')
            state = state[start:end+1]
            locations = {f'/{i}': f'JSON pointer {args.pointer}/{start+i}' for i in range(len(state))}
            read_scope += f'; selected zero-based indices {start}:{end}'
    if config.get('columns') and args.format not in ('csv', 'xlsx'):
        raise ValueError('columns mapping only applies to CSV/XLSX; supply normalized JSON or text')
    complete = {'yes':True, 'no':False, 'unknown':None}[args.complete]
    state, subjects = make_units(state, config, args.source_id, complete)
    if config['unit']['mode'] == 'group':
        mapped = {}
        for g, group in enumerate(state):
            for j, old_index in enumerate(group['input_indices']):
                prefix = f'/{old_index}'
                for path, physical in locations.items():
                    if path == prefix or path.startswith(prefix+'/'):
                        mapped[f'/{g}/rows/{j}'+path[len(prefix):]] = physical
                if not any(p.startswith(f'/{g}/rows/{j}') for p in mapped):
                    mapped[f'/{g}/rows/{j}'] = f'input array index {old_index}'
        locations = mapped
    metadata = {'filename': args.file.name}
    if config.get('columns'):
        metadata['column_mapping'] = config['columns']
    if args.format == 'xlsx':
        metadata.update(sheet=args.sheet, range=args.cell_range)
    if args.format == 'db-result' and args.query_id:
        metadata['query_id'] = args.query_id
    source = {'id':args.source_id, 'kind':args.format.replace('-', '_'), 'label':args.file.name,
              'scope':read_scope+'; intended scope: '+args.scope+'; completeness supplied by caller',
              'complete':complete, 'state_paths':[''], 'locations':locations, 'metadata':metadata}
    unit = {'mode':config['unit']['mode'], 'subjects':subjects, 'complete':complete}
    if unit['mode'] == 'group':
        unit['group_by'] = config['unit']['group_by']
    result = {'state':state, 'sources':[source], 'policy':config['policy'], 'questions':config['questions'], 'unit':unit}
    errors = validate_input(result)
    if errors:
        raise ValueError('; '.join(errors))
    return result


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--format', required=True, choices=['text','json','csv','xlsx','db-result','document'])
    p.add_argument('--file', type=Path, required=True)
    p.add_argument('--config', type=Path, required=True)
    p.add_argument('--source-id', required=True)
    p.add_argument('--scope', required=True, help='User-defined intended judgment population')
    p.add_argument('--complete', required=True, choices=['yes','no','unknown'], help='Caller-supplied completeness for that scope; never inferred')
    p.add_argument('--rows', help='CSV data records, 1-based inclusive, header excluded')
    p.add_argument('--lines', help='Text lines, 1-based inclusive')
    p.add_argument('--sheet')
    p.add_argument('--range', dest='cell_range', help='XLSX selected range including header')
    p.add_argument('--whole-file', action='store_true', help='Authorize bounded whole JSON file parsing')
    p.add_argument('--pointer', default='', help='JSON pointer selected after parsing')
    p.add_argument('--items', help='JSON array indices, zero-based inclusive')
    p.add_argument('--query-id', help='Existing query result ID, only if known')
    p.add_argument('--delimiter', default=',')
    p.add_argument('--encoding', default='utf-8-sig')
    args = p.parse_args()
    try:
        output = prepare(args)
    except (ValueError, TypeError, KeyError, IndexError, OSError, csv.Error, ImportError, zipfile.BadZipFile, OverflowError) as error:
        print(json.dumps({'stage':'preparation','ok':False,'error':str(error)}, ensure_ascii=False), file=sys.stderr)
        return 2
    print(json.dumps(output, ensure_ascii=False, indent=2, allow_nan=False))
    return 0


if __name__ == '__main__':
    sys.exit(main())
